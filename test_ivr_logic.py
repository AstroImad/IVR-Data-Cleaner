"""Regression tests for generalized IVR routing and skip-logic parsing."""

import io
import unittest
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document

from cleaning import (
    apply_flow_value_mapping,
    filter_skip_logic,
    load_csv_file,
    load_csvs_from_zip,
    validate_flow_values,
)
from parsers import get_skip_logic_candidates, parse_ivr_script


ROOT = Path(__file__).resolve().parent
EXAMPLES = ROOT / "example"


class CsvLoadingTests(unittest.TestCase):
    def test_fixed_width_report_still_loads(self):
        report = (
            "IVR report,,,,\n"
            "RequestTS,PhoneNo,VoiceMail,UserKeyPress,\n"
            "date,601111,,FlowNo_2=1,FlowNo_30=2\n"
        )
        loaded = load_csv_file(report.encode(), "old-format.csv")

        self.assertEqual(loaded.shape, (1, 3))
        self.assertEqual(loaded.loc[0, 0], "FlowNo_2=1")
        self.assertEqual(loaded.loc[0, 1], "FlowNo_30=2")

    def test_variable_width_report_preserves_answers_beyond_header(self):
        report = (
            "IVR report\n"
            "RequestTS,PhoneNo,LastStartAttemptDT,LastEndAttemptDT,"
            "LastAttemptReason,ConnectedDuration,VoiceMail,UserKeyPress\n"
            "date,601111,,,,,,FlowNo_2=1,FlowNo_30=2\n"
            "date,602222,,,,,,FlowNo_2=2\n"
        )
        loaded = load_csv_file(report.encode(), "new-format.csv")

        self.assertEqual(loaded.shape, (2, 3))
        self.assertEqual(loaded.loc[0, 0], "FlowNo_2=1")
        self.assertEqual(loaded.loc[0, 1], "FlowNo_30=2")
        self.assertTrue(pd.isna(loaded.loc[1, 1]))

    def test_zip_reports_failure_instead_of_returning_partial_data(self):
        good = (
            "IVR report\nRequestTS,PhoneNo,UserKeyPress\n"
            "date,601111,FlowNo_2=1\n"
        )
        bad = "IVR report\nRequestTS,PhoneNo,Other\ndate,602222,answer\n"
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zf:
            zf.writestr("good.csv", good)
            zf.writestr("bad.csv", bad)

        with self.assertRaisesRegex(ValueError, "bad.csv"):
            load_csvs_from_zip(archive.getvalue())


class IvrParserTests(unittest.TestCase):
    def test_all_example_documents_parse_with_routes(self):
        documents = sorted(EXAMPLES.glob("*.docx"))
        self.assertEqual(len(documents), 5)

        for document in documents:
            with self.subTest(document=document.name):
                questions, mappings, graph, branches = parse_ivr_script(
                    document.read_bytes(), document.name
                )
                self.assertTrue(questions)
                self.assertTrue(mappings)
                self.assertTrue(graph)
                self.assertIsInstance(branches, list)
                self.assertTrue(
                    any(info.get("routes") for info in graph.values()),
                    f"No routes detected in {document.name}",
                )

    def test_known_screening_routes_are_captured(self):
        expected = {
            "Johor.docx": (2, 2, 23),
            "Selangor Tracker State _ Julai 2026.docx": (2, 2, 29),
            "Script IVR Hulu Selangor.docx": (2, 2, 28),
        }
        for filename, (flow, choice, target) in expected.items():
            with self.subTest(filename=filename):
                questions, _, graph, _ = parse_ivr_script(
                    (EXAMPLES / filename).read_bytes(), filename
                )
                self.assertEqual(graph[flow]["answer_redirects"][choice], target)
                candidates = get_skip_logic_candidates(graph, questions)
                matching = [
                    candidate for candidate in candidates
                    if candidate["value"] == f"FlowNo_{flow}={choice}"
                ]
                self.assertTrue(matching)
                self.assertTrue(matching[0]["alternate"])

    def test_matrix_ranges_are_mapped_without_becoming_routes(self):
        filename = "Selangor Tracker State _ Julai 2026.docx"
        _, mappings, graph, _ = parse_ivr_script(
            (EXAMPLES / filename).read_bytes(), filename
        )

        self.assertEqual(mappings["FlowNo_3=1"], "Isu kesesakan lalu lintas")
        self.assertEqual(mappings["FlowNo_10=0"], "Anthony Loke")
        self.assertFalse(
            any(
                route["choice"] == 1
                for route in graph.get(3, {}).get("routes", [])
            )
        )

    def test_clsa_flow_13_has_exactly_seven_declared_answers(self):
        filename = "CLSA Soalan Survey IVR (1).docx"
        _, mappings, _, _ = parse_ivr_script(
            (EXAMPLES / filename).read_bytes(), filename
        )

        flow_13 = {
            key: value
            for key, value in mappings.items()
            if key.startswith("FlowNo_13=")
        }
        self.assertEqual(
            set(flow_13),
            {f"FlowNo_13={choice}" for choice in range(1, 8)},
        )

    def test_invalid_flow_13_singletons_remain_visible_and_unmapped(self):
        filename = "CLSA Soalan Survey IVR (1).docx"
        _, mappings, _, _ = parse_ivr_script(
            (EXAMPLES / filename).read_bytes(), filename
        )
        frame = pd.DataFrame(
            {
                "phonenum": ["valid", "multi", "zero", "eight", "nine"],
                "answer": [
                    "FlowNo_13=7",
                    "FlowNo_13=123",
                    "FlowNo_13=0",
                    "FlowNo_13=8",
                    "FlowNo_13=9",
                ],
            }
        )

        mapped = apply_flow_value_mapping(frame, mappings)
        self.assertEqual(mapped.loc[0, "answer"], "Lain-lain")
        self.assertEqual(
            mapped.loc[1, "answer"],
            "Barangan dapur; Makanan; Bil utiliti seperti elektrik dan air",
        )
        self.assertEqual(
            mapped.loc[2:, "answer"].tolist(),
            ["FlowNo_13=0", "FlowNo_13=8", "FlowNo_13=9"],
        )

        validation = validate_flow_values(frame, mappings)
        invalid = validation[validation["Raw Value"].isin(
            ["FlowNo_13=0", "FlowNo_13=8", "FlowNo_13=9"]
        )]
        self.assertEqual(invalid["Status"].tolist(), ["Unmapped"] * 3)
        self.assertEqual(invalid["Mapped Value"].tolist(), [""] * 3)

    def test_non_tekan_navigation_and_terminal_routes(self):
        document = Document()
        document.add_paragraph("Soalan saringan Call flow 2")
        document.add_paragraph("Option 1 -> flow 3")
        document.add_paragraph("Answer 2 untuk Tidak, pergi ke Call flow 9")
        document.add_paragraph("Soalan akhir Call flow 3")
        document.add_paragraph("Terima kasih Call flow 3")
        document.add_paragraph("Tamat Call flow 9")

        buffer = io.BytesIO()
        document.save(buffer)
        questions, _, graph, _ = parse_ivr_script(buffer.getvalue(), "synthetic.docx")

        self.assertEqual(graph[2]["answer_redirects"], {1: 3, 2: 9})
        self.assertTrue(graph[3]["terminal"])
        candidates = get_skip_logic_candidates(graph, questions)
        candidate = next(item for item in candidates if item["value"] == "FlowNo_2=2")
        self.assertTrue(candidate["terminal_reached"])

    def test_table_content_is_parsed(self):
        document = Document()
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Soalan saringan Call flow 2"
        table.cell(1, 0).text = "Tekan 2 untuk Tidak Call flow 9"

        buffer = io.BytesIO()
        document.save(buffer)
        _, _, graph, _ = parse_ivr_script(buffer.getvalue(), "table.docx")
        self.assertEqual(graph[2]["answer_redirects"][2], 9)


class SkipFilterTests(unittest.TestCase):
    def test_filter_accepts_multiple_values_and_multiselect_cells(self):
        frame = pd.DataFrame(
            {
                "phonenum": ["a", "b", "c", "d"],
                "answer": [
                    "FlowNo_2=1",
                    "FlowNo_2=2",
                    "FlowNo_3=4; FlowNo_2=2",
                    None,
                ],
            }
        )
        main, skipped = filter_skip_logic(frame, ["FlowNo_2=2", "FlowNo_3=4"])
        self.assertEqual(main["phonenum"].tolist(), ["a", "d"])
        self.assertEqual(skipped["phonenum"].tolist(), ["b", "c"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
